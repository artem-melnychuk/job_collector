"""Write an `active_near_fit` sheet: still-open postings worth applying to now.

This is the user's daily "where to apply" list, so it leaves out anything
already handled in the `applications` sheet (status sent or rejected),
matched by job_id or, for rows typed in by hand, by URL.

Joins manual_review (decision, should_be_filtered) with jobs_master
(availability_status, url, location) — the same join shape as
skills_gap_report.py — and keeps only postings that are both near-fit
(Подходит/Возможно, not filtered) and still live (availability_status ==
"active"). Point-in-time: availability_status only reflects the last
--check-availability run, not the current moment.
"""

from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core.ids import normalize_url
from core.storage import arrange_workbook_tabs

SHEET_NAME = "active_near_fit"
# `applied` first: the user marks "Да" there after sending; job_id last,
# needed to carry that mark over to manual_review.
COLUMNS = (
    "applied", "decision", "title", "company", "salary", "salary_usd_equivalent", "work_mode",
    "country", "city_region", "source", "also_on", "seniority_manual", "main_reason", "review_notes", "url",
    "job_id",
)
# Preferred copy when one job is listed several times (see core/cross_dedup.py):
# the company's own ATS first, then the boards with the most structured data.
SOURCE_PREFERENCE = ("Company Careers", "Djinni", "Himalayas", "Jobicy", "We Work Remotely", "RemoteOK", "LinkedIn")
APPLIED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
APPLIED_HEADER_FILL = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
# Statuses in the applications sheet that take a posting off the to-do list.
# "approved" (a draft ready to send) and "draft" stay on it.
HANDLED_STATUSES = {"sent", "rejected"}
DECISION_FILLS = {
    "Подходит": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "Возможно": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
}


def _sheet_rows(worksheet) -> list[dict[str, str]]:
    headers = [cell.value for cell in worksheet[1]]
    rows = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        rows.append({header: value for header, value in zip(headers, row) if header})
    return rows


def handled_postings(application_rows: list[dict[str, str]]) -> tuple[set[str], set[str]]:
    """job_ids and normalized URLs of postings already sent or rejected."""
    job_ids, urls = set(), set()
    for row in application_rows:
        if str(row.get("status") or "").strip() not in HANDLED_STATUSES:
            continue
        if row.get("job_id"):
            job_ids.add(str(row["job_id"]).strip())
        if row.get("url"):
            urls.add(normalize_url(str(row["url"])))
    return job_ids, urls


def _group(row: dict[str, str]) -> str:
    return str(row.get("duplicate_group") or "").strip()


def _source_rank(row: dict[str, str]) -> int:
    source = str(row.get("source") or "")
    return SOURCE_PREFERENCE.index(source) if source in SOURCE_PREFERENCE else len(SOURCE_PREFERENCE)


def build_active_near_fit(
    joined: list[dict[str, str]],
    handled_job_ids: set[str] = frozenset(),
    handled_urls: set[str] = frozenset(),
) -> list[dict[str, str]]:
    """Open near-fit postings not yet applied to, one row per job.

    Copies of the same job (same `duplicate_group`, see core/cross_dedup.py)
    count as one: applying to any copy takes the whole group off the list,
    and of the copies still eligible only one is shown, with the other
    copies' sources in `also_on`.
    """
    def handled(row: dict[str, str]) -> bool:
        return (
            str(row.get("job_id") or "") in handled_job_ids
            or normalize_url(str(row.get("url") or "")) in handled_urls
            or bool(str(row.get("applied") or "").strip())
        )

    handled_groups = {_group(row) for row in joined if _group(row) and handled(row)}
    kept = [
        row for row in joined
        if str(row.get("decision", "")).strip() in {"Подходит", "Возможно"}
        and str(row.get("should_be_filtered", "")).strip() != "Да"
        and str(row.get("availability_status", "")).strip() == "active"
        and not handled(row)
        and _group(row) not in handled_groups
    ]

    sources_by_group: dict[str, list[dict[str, str]]] = {}
    for row in joined:
        if _group(row):
            sources_by_group.setdefault(_group(row), []).append(row)
    best: dict[str, dict[str, str]] = {}
    for row in kept:
        group = _group(row)
        if not group:
            continue
        current = best.get(group)
        rank = (row.get("decision") != "Подходит", _source_rank(row), str(row.get("job_id") or ""))
        if current is None or rank < (current.get("decision") != "Подходит", _source_rank(current), str(current.get("job_id") or "")):
            best[group] = row
    collapsed = []
    for row in kept:
        group = _group(row)
        if group and best[group] is not row:
            continue
        if group:
            others = [other for other in sources_by_group[group] if other is not row]
            row = {**row, "also_on": ", ".join(sorted({str(other.get("source") or "") for other in others}))}
        collapsed.append(row)
    collapsed.sort(key=lambda row: (row.get("decision") != "Подходит", row.get("source", "")))
    return collapsed


def copy_applied_marks(workbook) -> int:
    """Move `applied` marks typed into this sheet over to manual_review.

    The sheet is rebuilt on every run, so a mark left only here would be
    lost; manual_review keeps it across reruns and is where
    scripts/applications_tracker.py reads marks from. Returns how many
    marks were copied.
    """
    if SHEET_NAME not in workbook.sheetnames or "manual_review" not in workbook.sheetnames:
        return 0
    marks = {
        str(row.get("job_id")): row.get("applied")
        for row in _sheet_rows(workbook[SHEET_NAME])
        if row.get("job_id") and str(row.get("applied") or "").strip()
    }
    review = workbook["manual_review"]
    headers = [cell.value for cell in review[1]]
    if not marks or "applied" not in headers or "job_id" not in headers:
        return 0
    job_id_col, applied_col = headers.index("job_id") + 1, headers.index("applied") + 1
    copied = 0
    for row_index in range(2, review.max_row + 1):
        job_id = str(review.cell(row_index, job_id_col).value or "")
        cell = review.cell(row_index, applied_col)
        if job_id in marks and not str(cell.value or "").strip():
            cell.value = marks[job_id]
            copied += 1
    return copied


def _write_sheet(workbook, rows: list[dict[str, str]]) -> None:
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    sheet = workbook.create_sheet(SHEET_NAME)
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    applied_col = COLUMNS.index("applied") + 1
    for row in rows:
        sheet.append([row.get(column, "") for column in COLUMNS])
        url_cell = sheet.cell(sheet.max_row, COLUMNS.index("url") + 1)
        if url_cell.value:
            url_cell.hyperlink = str(url_cell.value)
            url_cell.style = "Hyperlink"
        fill = DECISION_FILLS.get(row.get("decision"))
        if fill:
            for cell in sheet[sheet.max_row]:
                cell.fill = fill
        sheet.cell(sheet.max_row, applied_col).fill = APPLIED_FILL
    sheet.cell(1, applied_col).fill = APPLIED_HEADER_FILL
    applied_letter = sheet.cell(1, applied_col).column_letter
    validation = DataValidation(type="list", formula1='"Да"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"{applied_letter}2:{applied_letter}{max(sheet.max_row, 2)}")
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(max((len(str(cell.value or "")) for cell in column_cells), default=10) + 2, 50)
        sheet.column_dimensions[column_cells[0].column_letter].width = width


def refresh_active_near_fit(workbook) -> list[dict[str, str]]:
    """Rebuild the sheet in an already-open workbook (the caller saves)."""
    copied = copy_applied_marks(workbook)
    if copied:
        print(f"Copied {copied} 'applied' mark(s) from {SHEET_NAME} to manual_review")
    master_by_id = {row.get("job_id"): row for row in _sheet_rows(workbook["jobs_master"])}
    review_rows = _sheet_rows(workbook["manual_review"])

    joined = []
    for review in review_rows:
        master = master_by_id.get(review.get("job_id"), {})
        joined.append({**master, **review})

    application_rows = _sheet_rows(workbook["applications"]) if "applications" in workbook.sheetnames else []
    handled_job_ids, handled_urls = handled_postings(application_rows)
    rows = build_active_near_fit(joined, handled_job_ids, handled_urls)
    _write_sheet(workbook, rows)
    arrange_workbook_tabs(workbook)
    return rows


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    path = PROJECT_ROOT / "data" / "processed" / "crypto_jobs_clean_v1.xlsx"
    workbook = openpyxl.load_workbook(path)
    rows = refresh_active_near_fit(workbook)
    workbook.save(path)
    print(f"Active near-fit postings: {len(rows)}")
    for row in rows:
        print(f"  [{row['decision']}] {row['title']} — {row['company']} ({row['source']})")
    print()
    print(f"Saved '{SHEET_NAME}' sheet in {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
